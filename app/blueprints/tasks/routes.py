import os
import re
import json
from datetime import datetime
from urllib.parse import urlencode, quote
from flask import Blueprint, render_template, request, redirect, url_for, session, flash, jsonify, send_file
from werkzeug.utils import secure_filename
from sqlalchemy import func, text
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import (
    db, ActivityLog, ActivityLogComment, ActivityLogFile, 
    ProjectComment, User, Subsidiary, Tag, TagRequest, JiraHistory,
    Notification, AiConversation, JiraApiLog
)
from app.constants import Constants, JiraMapping
from app.utils.decorators import login_required, permission_required
from app.utils.helpers import status_activity_logs_class, assignee_class, priority_class, clean_html_comment
from app.services.excel_service import ExcelService
from app.services.jira_service import JiraService
from config import Config

tasks_bp = Blueprint('tasks', __name__)

def update_related_statuses(task):
    """Updates linked Tag and TagRequest statuses based on task status mapping, mirroring TasksController.php::updateRelatedStatuses."""
    if not task or task.status not in Constants.TASK_STATUS_REFERENCE_MAPPING:
        return

    # Do not update tags or tag-requests if this is a "New Tag Request" task that was just completed
    if task.status == Constants.ACTIVITY_LOGS_STATUS_COMPLETED and 'new tag request' in (task.title or '').lower():
        return

    mapping = Constants.TASK_STATUS_REFERENCE_MAPPING[task.status]

    # Update Tag Status
    if 'tag' in mapping and task.type in [Constants.ACTIVITY_TYPE_RENEWAL, Constants.ACTIVITY_TYPE_REQUEST] and task.team_id == Constants.TEAM_MEDIA_TAGGING:
        if task.ticket_id:
            tag_ids = []
            # 1. Try to get tags directly from Tag Requests
            requests_list = TagRequest.query.filter(
                (TagRequest.new_tag_ticket == task.ticket_id) | (TagRequest.new_rule_ticket == task.ticket_id)
            ).all()
            for req in requests_list:
                if req.tags:
                    for rid in str(req.tags).split(','):
                        rid_clean = rid.strip()
                        if rid_clean.isdigit():
                            tag_ids.append(int(rid_clean))

            # 2. Fallback to JiraHistory
            if not tag_ids:
                jira_tickets = JiraHistory.query.filter_by(ticket_id=task.ticket_id).all()
                for jt in jira_tickets:
                    if jt.tag_id:
                        tag_ids.append(int(jt.tag_id))
                    if jt.comments:
                        for rid in str(jt.comments).split(','):
                            rid_clean = rid.strip()
                            if rid_clean.isdigit():
                                tag_ids.append(int(rid_clean))

            tag_ids = list(set(tag_ids))
            if tag_ids:
                tags_query = Tag.query.filter(Tag.id.in_(tag_ids))
                if mapping['tag'] == Constants.STATUS_PENDING_VERIFICATION and task.type == Constants.ACTIVITY_TYPE_RENEWAL:
                    tags_query = tags_query.filter(Tag.requested_date.isnot(None), Tag.requested_date != '')
                tags_query.update({'status': mapping['tag']}, synchronize_session=False)
                db.session.commit()

    # Update Tag Request Status
    if 'request' in mapping and task.ticket_id:
        req_status = mapping['request']
        requests_list = TagRequest.query.filter(
            (TagRequest.new_tag_ticket == task.ticket_id) | (TagRequest.new_rule_ticket == task.ticket_id)
        ).all()
        for req in requests_list:
            if req.status != req_status:
                req.status = req_status
        db.session.commit()


import time

_AI_MODELS_CACHE = None
_AI_MODELS_CACHE_TIME = 0

def fetch_ai_models():
    """Fetches available LLM models from Fabrix API (cached 10 min) or returns supported defaults."""
    global _AI_MODELS_CACHE, _AI_MODELS_CACHE_TIME
    now = time.time()
    if _AI_MODELS_CACHE is not None and (now - _AI_MODELS_CACHE_TIME < 600):
        return _AI_MODELS_CACHE

    api_url = Config.FABRIX_API_URL.rstrip('/') + '/openapi/llm/v1/models'
    api_key = Config.FABRIX_API_TOKEN
    fabrix_client = Config.FABRIX_CLIENT
    headers = {
        'Accept': 'application/json',
        'x-fabrix-client': fabrix_client,
        'x-openapi-token': f"Bearer {api_key}"
    }
    models = None
    if api_key and fabrix_client:
        try:
            resp = requests.get(api_url, headers=headers, timeout=2)
            if 200 <= resp.status_code < 300:
                models = resp.json().get('data', [])
        except Exception:
            pass

    if not models:
        models = [
            {'id': 'gpt-4o', 'name': 'GPT-4o'},
            {'id': 'gpt-4o-mini', 'name': 'GPT-4o mini'}
        ]
    _AI_MODELS_CACHE = models
    _AI_MODELS_CACHE_TIME = now
    return _AI_MODELS_CACHE


def get_user_color(user_id):
    color_map = {
        Constants.USER_MARIE: '#7B1FA2',
        Constants.USER_KHRISTINE: '#C2185B',
        Constants.USER_LUCKY: '#1B5E20',
        Constants.USER_JANITHA: '#1565C0',
        Constants.USER_GEORGE: '#E65100',
    }
    base = color_map.get(user_id, '#0071E3')
    hex_val = base.lstrip('#')
    return {
        'user_color': base,
        'user_color_r': int(hex_val[0:2], 16),
        'user_color_g': int(hex_val[2:4], 16),
        'user_color_b': int(hex_val[4:6], 16),
    }

@tasks_bp.route('/')
@login_required
@permission_required('tasks')
def index():
    page = request.args.get('page', 1, type=int)
    per_page = 25

    # Me mode handling
    me_mode_get = request.args.get('me_mode')
    if me_mode_get is not None:
        session['al_me_mode'] = 1 if me_mode_get == '1' else 0
    me_mode = (session.get('al_me_mode') == 1)

    filters = {
        'keyword': request.args.get('keyword', '').strip(),
        'subsidiary_id': request.args.get('subsidiary_id'),
        'team_id': request.args.get('team_id'),
        'project': request.args.get('project', '').strip(),
        'type': request.args.get('type'),
        'status': request.args.get('status'),
        'status_group': request.args.get('status_group'),
        'priority': request.args.get('priority'),
        'assignee': request.args.get('assignee'),
        'me_mode': 1 if me_mode else 0,
        'sort_by': request.args.get('sort_by') or 'created_at',
        'sort_dir': request.args.get('sort_dir') or 'desc',
        'show_completed': request.args.get('show_completed'),
        'created_from': request.args.get('created_from'),
        'created_to': request.args.get('created_to'),
        'started_from': request.args.get('started_from'),
        'started_to': request.args.get('started_to'),
        'finished_from': request.args.get('finished_from'),
        'finished_to': request.args.get('finished_to'),
        'labels': request.args.getlist('labels') or request.args.get('labels'),
    }

    advanced_filter_keys = [
        'subsidiary_id', 'team_id', 'project', 'assignee', 'status',
        'priority', 'type', 'labels', 'created_from', 'created_to',
        'started_from', 'started_to', 'finished_from', 'finished_to', 'campaign_status'
    ]
    advanced_filters_count = sum(1 for k in advanced_filter_keys if filters.get(k))
    has_active_advanced_filters = (advanced_filters_count > 0)

    from urllib.parse import urlencode

    status_counts = ActivityLog.get_status_counts(filters)
    needs_attention_count = ActivityLog.get_needs_attention_count(filters)
    pagination = ActivityLog.filter_logs(filters, paginate=per_page, page=page)
    logs = pagination.items

    pending_count = (status_counts.get(1, 0)) + (status_counts.get(5, 0))
    in_progress_count = (status_counts.get(2, 0)) + (status_counts.get(3, 0)) + (status_counts.get(4, 0))
    blocked_count = status_counts.get(7, 0)
    completed_count = status_counts.get(6, 0)
    is_me_mode = (request.args.get('me_mode') == '1') if 'me_mode' in request.args else (session.get('al_me_mode') == 1)

    # Preload user map & subsidiary map for labels
    all_users_map = {u.id: u.name for u in User.query.all()}
    all_subs_map = {s.id: s.name for s in Subsidiary.query.all()}

    active_pills = []
    all_filter_keys = [
        'subsidiary_id', 'team_id', 'project', 'assignee', 'status',
        'priority', 'type', 'labels', 'created_from', 'created_to',
        'started_from', 'started_to', 'finished_from', 'finished_to',
        'status_group', 'campaign_status', 'show_completed', 'keyword'
    ]

    for key in all_filter_keys:
        val = request.args.get(key)
        if not val or val == '':
            continue
        label = ''
        if key == 'subsidiary_id':
            sub_name = all_subs_map.get(int(val), val) if str(val).isdigit() else val
            label = f"Subsidiary: {sub_name}"
        elif key == 'team_id':
            team_name = Constants.TEAMS.get(int(val), val) if str(val).isdigit() else val
            label = f"Team: {team_name}"
        elif key == 'project':
            label = f"Project: {val}"
        elif key == 'assignee':
            u_name = all_users_map.get(int(val), val) if str(val).isdigit() else val
            label = f"Assignee: {u_name}"
        elif key == 'status':
            st_name = Constants.ACTIVITY_LOGS_STATUS.get(int(val), val) if str(val).isdigit() else val
            label = f"Status: {st_name}"
        elif key == 'status_group':
            group_name = val.replace('_', ' ').title()
            label = f"Group: {group_name}"
        elif key == 'priority':
            prio_name = Constants.ACTIVITY_PRIORITY.get(int(val), val) if str(val).isdigit() else val
            label = f"Priority: {prio_name}"
        elif key == 'type':
            type_name = Constants.ACTIVITY_TYPES.get(int(val), val) if str(val).isdigit() else val
            label = f"Type: {type_name}"
        elif key == 'labels':
            label = f"Labels: {val}"
        elif key == 'created_from':
            label = f"Created From: {val}"
        elif key == 'created_to':
            label = f"Created To: {val}"
        elif key == 'started_from':
            label = f"Started From: {val}"
        elif key == 'started_to':
            label = f"Started To: {val}"
        elif key == 'finished_from':
            label = f"Finished From: {val}"
        elif key == 'finished_to':
            label = f"Finished To: {val}"
        elif key == 'campaign_status':
            cs_name = Constants.CAMPAIGN_STATUS.get(int(val), val) if str(val).isdigit() else val
            label = f"Campaign: {cs_name}"
        elif key == 'show_completed':
            if val == '1':
                label = "Show Completed: Yes"
        elif key == 'keyword':
            label = f"Search: {val}"

        if label:
            args_copy = request.args.to_dict()
            args_copy.pop(key, None)
            args_copy.pop('page', None)
            qs = urlencode(args_copy)
            clear_url = request.path + ('?' + qs if qs else '')
            active_pills.append({
                'key': key,
                'label': label,
                'clear_url': clear_url
            })

    project_names = ActivityLog.get_distinct_projects()
    all_labels = ActivityLog.get_all_labels()
    colors = get_user_color(session.get('user_id'))
    user_name = session.get('user_name', 'there')
    welcome_message = f"Hello {user_name}! Do you have any questions regarding the tasks? 😄"

    return render_template(
        'tasks/list.html',
        logs=logs,
        pagination=pagination,
        status_counts=status_counts,
        needs_attention_count=needs_attention_count,
        pendingCount=pending_count,
        inProgressCount=in_progress_count,
        blockedCount=blocked_count,
        completedCount=completed_count,
        isMeMode=is_me_mode,
        totalLogs=pagination.total,
        active_pills=active_pills,
        advancedFiltersCount=advanced_filters_count,
        hasActiveAdvancedFilters=has_active_advanced_filters,
        activeFiltersCount=advanced_filters_count,
        project_names=project_names,
        all_labels=all_labels,
        filters=filters,
        chatWelcomeMessage=welcome_message,
        aiModels=fetch_ai_models(),
        **colors
    )

@tasks_bp.route('/<int:id>', endpoint='view_short')
@tasks_bp.route('/view/<int:id>')
@login_required
@permission_required('tasks')
def view(id):
    task = ActivityLog.query.get_or_404(id)
    comments = ActivityLogComment.query.filter_by(log_id=task.id).order_by(ActivityLogComment.created_at.asc()).all()
    files = ActivityLogFile.query.filter_by(log_id=task.id).order_by(ActivityLogFile.created_at.desc()).all()
    all_users = User.query.filter_by(is_active=1).all()

    # Query related tags if type is Renewal/Request and Team is Media Tagging, mirroring TasksController.php::view
    show_related_tags = (
        (task.type in [Constants.ACTIVITY_TYPE_RENEWAL, Constants.ACTIVITY_TYPE_REQUEST])
        and task.team_id == Constants.TEAM_MEDIA_TAGGING
    )
    related_tags = []
    jira_ticket_record = None

    if show_related_tags and task.ticket_id:
        jira_tickets = JiraHistory.query.filter_by(ticket_id=task.ticket_id).all()
        if jira_tickets:
            jira_ticket_record = jira_tickets[0]
            tag_ids = []
            for jt in jira_tickets:
                if jt.tag_id:
                    tag_ids.append(int(jt.tag_id))
                if jt.comments:
                    for rid in str(jt.comments).split(','):
                        rid_clean = rid.strip()
                        if rid_clean.isdigit():
                            tag_ids.append(int(rid_clean))
            tag_ids = list(set(tag_ids))
            if tag_ids:
                related_tags = Tag.query.filter(Tag.id.in_(tag_ids), Tag.deleted_at.is_(None)).all()

    colors = get_user_color(session.get('user_id'))

    return render_template(
        'tasks/view.html',
        task=task,
        log=task,
        comments=comments,
        files=files,
        all_users=all_users,
        showRelatedTags=show_related_tags,
        relatedTags=related_tags,
        jiraTicketRecord=jira_ticket_record,
        **colors
    )

@tasks_bp.route('/project-spreadsheet')
@tasks_bp.route('/project')
@tasks_bp.route('/project/')
@login_required
@permission_required('tasks')
def project_spreadsheet_default():
    show_completed = (request.args.get('show_completed') == '1')

    # Get active projects (projects that have at least one non-completed task)
    active_rows = db.session.query(ActivityLog.project).filter(
        ActivityLog.project.isnot(None),
        ActivityLog.project != '',
        ActivityLog.deleted_at.is_(None),
        ActivityLog.status != Constants.ACTIVITY_LOGS_STATUS_COMPLETED
    ).distinct().order_by(ActivityLog.project).all()
    active_projects = [r[0] for r in active_rows if r[0]]

    all_projects = ActivityLog.get_distinct_projects()

    if show_completed and all_projects:
        first_proj = all_projects[0]
    elif active_projects:
        first_proj = active_projects[0]
    elif all_projects:
        first_proj = all_projects[0]
    else:
        first_proj = 'Generic'

    args = request.args.to_dict()
    if str(args.get('show_completed', '')) in ['0', '']:
        args.pop('show_completed', None)
    query_str = ('?' + urlencode(args)) if args else ''
    return redirect(f"/tasks/project/{quote(str(first_proj))}{query_str}")

@tasks_bp.route('/project/<path:project_name>')
@login_required
@permission_required('tasks')
def project_view(project_name):

    page = request.args.get('page', 1, type=int)
    per_page = 25

    me_mode_get = request.args.get('me_mode')
    if me_mode_get is not None:
        session['al_me_mode'] = 1 if me_mode_get == '1' else 0
    me_mode = (session.get('al_me_mode') == 1)

    show_completed = (request.args.get('show_completed') == '1')

    tasks = ActivityLog.get_project_logs(
        project_name=project_name,
        me_mode=me_mode,
        show_completed=show_completed
    )

    # Active projects (projects with at least one non-completed, non-deleted task)
    active_rows = db.session.query(ActivityLog.project).filter(
        ActivityLog.project.isnot(None),
        ActivityLog.project != '',
        ActivityLog.deleted_at.is_(None),
        ActivityLog.status != Constants.ACTIVITY_LOGS_STATUS_COMPLETED
    ).distinct().all()
    active_projects = set(r[0] for r in active_rows if r[0])

    all_projects = ActivityLog.get_distinct_projects()
    all_project_names = [p if p else 'Generic' for p in all_projects]
    all_project_names = sorted(list(set(all_project_names)))

    if show_completed:
        # Show all project tabs including completed projects
        project_names = all_project_names
        if project_name not in project_names:
            project_names.append(project_name)
            project_names.sort()
    else:
        # Completed project tabs should NOT display unless Show Completed is checked
        project_names = [p for p in all_project_names if p in active_projects]
        if not project_names:
            project_names = [project_name] if project_name else ['Generic']

    # Calculate project progress
    total_proj = ActivityLog.query.filter_by(project=project_name, deleted_at=None).count()
    completed_proj = ActivityLog.query.filter_by(project=project_name, status=Constants.ACTIVITY_LOGS_STATUS_COMPLETED, deleted_at=None).count()
    progress_percentage = int((completed_proj / total_proj) * 100) if total_proj > 0 else 0

    # Project comment counts
    comment_counts_query = db.session.query(
        ProjectComment.project_name, func.count(ProjectComment.id)
    ).filter(ProjectComment.deleted_at.is_(None)).group_by(ProjectComment.project_name).all()
    project_comment_counts = {r[0] if r[0] else 'Generic': r[1] for r in comment_counts_query}

    colors = get_user_color(session.get('user_id'))

    return render_template(
        'tasks/project_spreadsheet.html',
        project_name=project_name,
        tasks=tasks,
        project_names=project_names,
        me_mode=me_mode,
        show_completed=show_completed,
        progress_percentage=progress_percentage,
        project_comment_counts=project_comment_counts,
        **colors
    )

@tasks_bp.route('/create', methods=['GET', 'POST'])
@login_required
@permission_required('tasks')
def create():
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        description = request.form.get('description', '').strip()
        project = request.form.get('project', '').strip() or 'Generic'
        ticket_id = request.form.get('ticket_id', '').strip()
        assignee = request.form.get('assignee', type=int) or session.get('user_id')
        requested_by = request.form.get('requested_by', '').strip()
        subsidiary_id = request.form.get('subsidiary_id', type=int) or 1
        team_id = request.form.get('team_id', type=int) or Constants.USER_TEAMS.get(session.get('user_id'), Constants.TEAM_MEDIA_TAGGING)
        department_id = request.form.get('department_id', '').strip()
        campaign_status = request.form.get('campaign_status', type=int)
        status = request.form.get('status', type=int) or Constants.ACTIVITY_LOGS_STATUS_TO_DO
        priority = request.form.get('priority', type=int) or Constants.PRIORITY_MEDIUM
        size = request.form.get('size', type=int) or Constants.ACTIVITY_SIZE_M
        activity_type = request.form.get('type', type=int) or Constants.ACTIVITY_TYPE_TASK

        labels_list = request.form.getlist('labels[]') or request.form.getlist('labels')
        labels_str = ','.join([l.strip().replace(' ', '') for l in labels_list if l.strip()])

        req_date = request.form.get('requested_date')
        start_date = request.form.get('start_date')
        finish_date = request.form.get('finished_date')

        task = ActivityLog(
            title=title,
            description=description,
            project=project,
            ticket_id=ticket_id,
            team_id=team_id,
            assignee=assignee,
            requested_by=requested_by,
            subsidiary_id=subsidiary_id,
            department_id=department_id,
            campaign_status=campaign_status,
            status=status,
            priority=priority,
            size=size,
            type=activity_type,
            labels=labels_str,
            requested_date=datetime.strptime(req_date, '%Y-%m-%d') if req_date else None,
            start_date=datetime.strptime(start_date, '%Y-%m-%d') if start_date else None,
            finished_date=datetime.strptime(finish_date, '%Y-%m-%d') if finish_date else None,
            created_by=session.get('user_id', 0)
        )
        db.session.add(task)
        db.session.commit()

        # Notify assignee if assigned to someone else
        if assignee and assignee in Constants.USERS and assignee != session.get('user_id'):
            notif = Notification(
                user_id=int(assignee),
                sender_id=session.get('user_id', 0),
                task_id=task.id,
                message=f"assigned the task '{task.title}' to you.",
                is_read=0
            )
            db.session.add(notif)
            db.session.commit()

        flash('Task created successfully.', 'success')

        referer = request.headers.get('Referer')
        if referer and '/tasks/project/' in referer:
            return redirect(url_for('tasks.project_view', project_name=task.project))
        return redirect(url_for('tasks.view', id=task.id))

    # Duplicate pre-fill or project pre-fill
    log = None
    duplicate_id = request.args.get('duplicate_id', type=int)
    if duplicate_id:
        orig = ActivityLog.query.get(duplicate_id)
        if orig:
            log = orig.to_dict() if hasattr(orig, 'to_dict') else orig
            log['title'] = f"Copy of {orig.title}"
            log['finished_date'] = None

    proj_arg = request.args.get('project', '').strip()
    if proj_arg:
        if log is None:
            log = {}
        log['project'] = proj_arg

    project_names = ActivityLog.get_distinct_projects()
    all_labels = ActivityLog.get_all_labels()
    departments = [r[0] for r in db.session.query(ActivityLog.department_id).distinct().filter(ActivityLog.department_id.isnot(None), ActivityLog.department_id != '').all()]

    return render_template(
        'tasks/create.html',
        log=log,
        project_names=project_names,
        department_names=departments,
        all_labels=all_labels
    )

@tasks_bp.route('/update/<int:id>', methods=['GET', 'POST'])
@login_required
@permission_required('tasks')
def update(id):
    task = ActivityLog.query.get_or_404(id)
    if request.method == 'POST':
        old_status = task.status
        task.title = request.form.get('title', '').strip()
        task.description = request.form.get('description', '').strip()
        task.project = request.form.get('project', '').strip() or 'Generic'
        task.ticket_id = request.form.get('ticket_id', '').strip() or None
        task.mt_jira_ticket = request.form.get('mt_jira_ticket', '').strip() or None
        task.team_id = request.form.get('team_id', type=int) or task.team_id
        task.assignee = request.form.get('assignee', type=int) or task.assignee
        task.requested_by = request.form.get('requested_by', '').strip() or None
        task.subsidiary_id = request.form.get('subsidiary_id', type=int) or task.subsidiary_id
        task.department_id = request.form.get('department_id', '').strip() or None
        task.campaign_status = request.form.get('campaign_status', type=int)
        task.status = request.form.get('status', type=int) or task.status
        task.priority = request.form.get('priority', type=int) or task.priority
        task.size = request.form.get('size', type=int) or task.size
        task.type = request.form.get('type', type=int) or task.type

        labels_list = request.form.getlist('labels[]') or request.form.getlist('labels')
        if labels_list:
            task.labels = ','.join([l.strip().replace(' ', '') for l in labels_list if l.strip()])

        req_date = request.form.get('requested_date')
        if req_date:
            task.requested_date = datetime.strptime(req_date, '%Y-%m-%d')
        start_date = request.form.get('start_date')
        if start_date:
            task.start_date = datetime.strptime(start_date, '%Y-%m-%d')
        finish_date = request.form.get('finished_date')
        if finish_date:
            task.finished_date = datetime.strptime(finish_date, '%Y-%m-%d')

        if task.status == Constants.ACTIVITY_LOGS_STATUS_COMPLETED and not task.finished_date:
            task.finished_date = datetime.utcnow()

        task.updated_by = session.get('user_id', 0)
        db.session.commit()

        # Update linked Tag & TagRequest statuses
        update_related_statuses(task)

        # Trigger JIRA issue update
        JiraService.edit_issue(task.id)

        if old_status != task.status:
            JiraService.update_status(task.id, old_status, task.status)

        flash('Task updated successfully.', 'success')
        return redirect(url_for('tasks.view', id=task.id))

    project_names = ActivityLog.get_distinct_projects()
    all_labels = ActivityLog.get_all_labels()
    departments = [r[0] for r in db.session.query(ActivityLog.department_id).distinct().filter(ActivityLog.department_id.isnot(None), ActivityLog.department_id != '').all()]

    return render_template(
        'tasks/update.html',
        log=task,
        project_names=project_names,
        department_names=departments,
        all_labels=all_labels
    )

@tasks_bp.route('/update-status/<int:id>', methods=['POST'])
@tasks_bp.route('/update-status', methods=['POST'])
@login_required
def update_status(id=None):
    task_id = id or request.form.get('id', type=int) or request.form.get('task_id', type=int)
    task = ActivityLog.query.get(task_id)
    if not task:
        return jsonify({'success': False, 'message': 'Task not found.'}), 404

    new_status = request.form.get('status', type=int)
    if new_status and new_status in Constants.ACTIVITY_LOGS_STATUS:
        old_status = task.status
        old_name = Constants.ACTIVITY_LOGS_STATUS.get(old_status, 'Unknown')
        to_name = Constants.ACTIVITY_LOGS_STATUS.get(new_status, 'Unknown')
        task.status = new_status
        task.updated_by = session.get('user_id', 0)
        if new_status == Constants.ACTIVITY_LOGS_STATUS_COMPLETED:
            task.finished_date = datetime.utcnow()

        comment = ActivityLogComment(
            log_id=task.id,
            comment=f"Status updated from [{old_name}] to [{to_name}]",
            type=Constants.COMMENT_TYPE_UPDATE,
            commented_by=session.get('user_id', 0),
            created_at=datetime.utcnow()
        )
        db.session.add(comment)
        db.session.commit()

        # Update linked Tag & TagRequest statuses
        update_related_statuses(task)

        # Sync transition to Jira if ticket exists
        JiraService.update_status(task.id, old_status, new_status)

        return jsonify({
            'success': True,
            'message': f"Status updated to '{to_name}' successfully.",
            'status_class': status_activity_logs_class(new_status),
            'comment': {
                'comment': comment.comment,
                'commented_by': comment.commented_by,
                'created_at': comment.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                'user_label': Constants.USERS.get(comment.commented_by, Constants.USER_SYSTEM_LABEL),
                'is_system': comment.commented_by == Constants.USER_SYSTEM
            }
        })
    return jsonify({'success': False, 'message': 'Invalid status'}), 400

@tasks_bp.route('/update-assignee/<int:id>', methods=['POST'])
@tasks_bp.route('/update-assignee', methods=['POST'])
@login_required
def update_assignee(id=None):
    task_id = id or request.form.get('id', type=int) or request.form.get('task_id', type=int)
    task = ActivityLog.query.get(task_id)
    if not task:
        return jsonify({'success': False, 'message': 'Task not found.'}), 404

    new_assignee = request.form.get('assignee', type=int)
    if new_assignee is not None:
        old_assignee = task.assignee
        old_name = Constants.USERS.get(old_assignee, Constants.USER_SYSTEM_LABEL)
        to_name = Constants.USERS.get(new_assignee, Constants.USER_SYSTEM_LABEL)
        task.assignee = new_assignee
        task.updated_by = session.get('user_id', 0)

        comment = ActivityLogComment(
            log_id=task.id,
            comment=f"Assignee changed from [{old_name}] to [{to_name}]",
            type=Constants.COMMENT_TYPE_UPDATE,
            commented_by=session.get('user_id', 0),
            created_at=datetime.utcnow()
        )
        db.session.add(comment)

        # Notify assigned user if different
        if new_assignee and new_assignee != session.get('user_id') and new_assignee in Constants.USERS:
            notif = Notification(
                user_id=int(new_assignee),
                sender_id=session.get('user_id', 0),
                task_id=task.id,
                message=f"assigned the task '{task.title}' to you.",
                is_read=0
            )
            db.session.add(notif)

        db.session.commit()

        # Update Jira issue assignee
        JiraService.edit_issue(task.id)

        return jsonify({
            'success': True,
            'message': f"Assignee updated to '{to_name}' successfully.",
            'assignee_class': assignee_class(new_assignee)
        })
    return jsonify({'success': False, 'message': 'Invalid assignee'}), 400

@tasks_bp.route('/update-priority/<int:id>', methods=['POST'])
@tasks_bp.route('/update-priority', methods=['POST'])
@login_required
def update_priority(id=None):
    task_id = id or request.form.get('id', type=int) or request.form.get('task_id', type=int)
    task = ActivityLog.query.get(task_id)
    if not task:
        return jsonify({'success': False, 'message': 'Task not found.'}), 404

    new_pri = request.form.get('priority', type=int)
    if new_pri and new_pri in Constants.PRIORITIES:
        to_name = Constants.PRIORITIES.get(new_pri, 'Normal')
        task.priority = new_pri
        task.updated_by = session.get('user_id', 0)
        db.session.commit()

        # Update Jira issue priority
        JiraService.edit_issue(task.id)

        return jsonify({
            'success': True,
            'message': f"Priority updated to '{to_name}'.",
            'priority_class': priority_class(new_pri)
        })
    return jsonify({'success': False, 'message': 'Invalid priority'}), 400

@tasks_bp.route('/update-ticket-id/<int:id>', methods=['POST'])
@tasks_bp.route('/update-ticket-id', methods=['POST'])
@login_required
def update_ticket_id(id=None):
    task_id = id or request.form.get('id', type=int) or request.form.get('task_id', type=int)
    task = ActivityLog.query.get(task_id)
    if not task:
        return jsonify({'success': False, 'message': 'Task not found.'}), 404

    t_id = request.form.get('ticket_id', '').strip()
    task.ticket_id = t_id
    task.updated_by = session.get('user_id', 0)
    db.session.commit()
    return jsonify({'success': True, 'ticket_id': t_id, 'jira_url': f"{Constants.JIRA_LINK}{t_id}" if t_id else ''})

@tasks_bp.route('/create-jira/<int:id>', methods=['POST'])
@tasks_bp.route('/create-jira-issue/<int:id>', methods=['POST'])
@login_required
def create_jira_issue(id):
    res = JiraService.create_issue(id)
    if res.get('success'):
        task = ActivityLog.query.get(id)
        ticket_id = task.mt_jira_ticket if task else res.get('ticket_id', '')
        return jsonify({
            'success': True,
            'message': 'MT JIRA Ticket created successfully.',
            'ticket_id': ticket_id,
            'csrf_hash': ''
        })
    return jsonify({
        'success': False,
        'message': res.get('message', 'Failed to create MT JIRA Ticket. Please check API Logs.'),
        'csrf_hash': ''
    })

@tasks_bp.route('/remove-tag/<int:id>', methods=['POST'])
@tasks_bp.route('/remove-tag', methods=['POST'])
@login_required
def remove_tag(id=None):
    task_id = id or request.form.get('id', type=int) or request.form.get('task_id', type=int)
    task = ActivityLog.query.get(task_id)
    if not task:
        return jsonify({'success': False, 'message': 'Activity log not found.', 'csrf_hash': ''})

    tag_id_to_remove = request.form.get('tag_id', type=int)
    if not tag_id_to_remove:
        return jsonify({'success': False, 'message': 'Invalid tag ID.', 'csrf_hash': ''})

    if not task.ticket_id:
        return jsonify({'success': False, 'message': 'Task is not associated with any JIRA ticket ID.', 'csrf_hash': ''})

    jira_tickets = JiraHistory.query.filter_by(ticket_id=task.ticket_id).all()
    if not jira_tickets:
        return jsonify({'success': False, 'message': 'No JIRA history entries found for this task.', 'csrf_hash': ''})

    for jt in jira_tickets:
        comments_tag_ids = []
        if jt.comments:
            for rid in str(jt.comments).split(','):
                rid_clean = rid.strip()
                if rid_clean.isdigit():
                    comments_tag_ids.append(int(rid_clean))

        new_comments_tag_ids = [tid for tid in comments_tag_ids if tid != tag_id_to_remove]
        new_tag_id = jt.tag_id
        if jt.tag_id == tag_id_to_remove:
            new_tag_id = new_comments_tag_ids[0] if new_comments_tag_ids else None

        jt.comments = ','.join(map(str, new_comments_tag_ids))
        jt.tag_id = new_tag_id
        jt.updated_by = session.get('user_id', 0)

    tag = Tag.query.get(tag_id_to_remove)
    tag_name = tag.name if tag else f"#{tag_id_to_remove}"

    comment = ActivityLogComment(
        log_id=task.id,
        comment=f"Removed tag [{tag_name}] from this task.",
        commented_by=session.get('user_id', Constants.USER_SYSTEM),
        type=Constants.COMMENT_TYPE_UPDATE,
        created_at=datetime.utcnow()
    )
    db.session.add(comment)
    db.session.commit()

    return jsonify({
        'success': True,
        'message': f"Tag '{tag_name}' removed from task.",
        'csrf_hash': '',
        'comment': {
            'comment': comment.comment,
            'commented_by': comment.commented_by,
            'created_at': comment.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'user_label': Constants.USERS.get(comment.commented_by, Constants.USER_SYSTEM_LABEL),
            'is_system': comment.commented_by == Constants.USER_SYSTEM
        }
    })


@tasks_bp.route('/bulk-status', methods=['POST'])
@login_required
def bulk_status():
    log_ids = request.form.getlist('log_ids[]') or request.form.getlist('log_ids')
    new_status = request.form.get('status', type=int)
    if log_ids and new_status:
        for lid in log_ids:
            task = ActivityLog.query.get(int(lid))
            if task:
                task.status = new_status
                task.updated_by = session.get('user_id', 0)
        db.session.commit()
        return jsonify({'success': True, 'message': f"Updated {len(log_ids)} tasks."})
    return jsonify({'success': False, 'message': 'Missing parameters.'}), 400

@tasks_bp.route('/bulk-project', methods=['POST'])
@login_required
def bulk_project():
    log_ids = request.form.getlist('log_ids[]') or request.form.getlist('log_ids')
    new_proj = request.form.get('project', '').strip()
    if log_ids and new_proj:
        for lid in log_ids:
            task = ActivityLog.query.get(int(lid))
            if task:
                task.project = new_proj
                task.updated_by = session.get('user_id', 0)
        db.session.commit()
        return jsonify({'success': True, 'message': f"Assigned {len(log_ids)} tasks to '{new_proj}'."})
    return jsonify({'success': False, 'message': 'Missing parameters.'}), 400

@tasks_bp.route('/bulk-assignee', methods=['POST'])
@login_required
def bulk_assignee():
    log_ids = request.form.getlist('log_ids[]') or request.form.getlist('log_ids')
    new_assignee = request.form.get('assignee', type=int)
    if log_ids and new_assignee is not None:
        for lid in log_ids:
            task = ActivityLog.query.get(int(lid))
            if task:
                task.assignee = new_assignee
                task.updated_by = session.get('user_id', 0)
        db.session.commit()
        return jsonify({'success': True, 'message': f"Assigned {len(log_ids)} tasks."})
    return jsonify({'success': False, 'message': 'Missing parameters.'}), 400

@tasks_bp.route('/comment', methods=['POST'])
@tasks_bp.route('/tasks/comment', methods=['POST'])
@login_required
def post_comment():
    task_id = request.form.get('task_id', type=int) or request.form.get('log_id', type=int)
    raw_comment = request.form.get('comment', '').strip()
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest' or request.is_json

    if not task_id or not raw_comment:
        if is_ajax:
            return jsonify({'success': False, 'message': 'No comment data provided.', 'csrf_hash': ''}), 400
        flash('No comment data provided.', 'error')
        return redirect(request.referrer or url_for('tasks.index'))

    clean_comment = clean_html_comment(raw_comment)
    current_uid = session.get('user_id', Constants.USER_SYSTEM)

    comment = ActivityLogComment(
        log_id=task_id,
        comment=clean_comment,
        type=Constants.COMMENT_TYPE_COMMENT,
        commented_by=current_uid,
        created_at=datetime.utcnow()
    )
    db.session.add(comment)
    db.session.commit()

    # Sync comment to Jira if task has mt_jira_ticket
    task = ActivityLog.query.get(task_id)
    if task and task.mt_jira_ticket:
        commenter_name = Constants.USERS.get(current_uid, session.get('user_name', 'System'))
        extracted_images = []
        formatted_jira_body = JiraService.format_html_for_jira(clean_comment, extracted_images)
        jira_comment_text = f"{commenter_name} commented:\n{formatted_jira_body}"
        JiraService.add_comment(task.id, task.mt_jira_ticket, jira_comment_text)

        # Upload inline images if any were extracted
        if extracted_images:
            for img in extracted_images:
                JiraService.upload_attachment(task.id, task.mt_jira_ticket, img['path'], img['name'])
                try:
                    os.unlink(img['path'])
                except Exception:
                    pass

    # Process @mentions
    mentioned_names = set(re.findall(r'@(\w+)', clean_comment))
    for name in mentioned_names:
        matched_user_id = None
        for uid, uname in Constants.USERS.items():
            if uname.lower() == name.lower():
                matched_user_id = uid
                break

        if matched_user_id and matched_user_id != current_uid:
            notif = Notification(
                user_id=matched_user_id,
                sender_id=current_uid,
                task_id=task_id,
                message='mentioned you in an update.',
                is_read=0,
                created_at=datetime.utcnow()
            )
            db.session.add(notif)
    db.session.commit()

    if is_ajax:
        return jsonify({
            'success': True,
            'message': 'Update posted successfully.',
            'csrf_hash': '',
            'comment': {
                'id': comment.id,
                'comment': comment.comment,
                'commented_by': comment.commented_by,
                'created_at': comment.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                'user_label': Constants.USERS.get(comment.commented_by, Constants.USER_SYSTEM_LABEL),
                'is_system': comment.commented_by == Constants.USER_SYSTEM
            }
        })

    flash('Update posted successfully.', 'success')
    return redirect(url_for('tasks.view', id=task_id))


@tasks_bp.route('/upload', methods=['POST'])
@tasks_bp.route('/tasks/upload', methods=['POST'])
@login_required
def upload_file():
    task_id = request.form.get('log_id', type=int) or request.form.get('task_id', type=int)
    file = request.files.get('file') or request.files.get('file_path')

    if not task_id or not file or not file.filename:
        flash('An error occurred. File was not uploaded successfully.', 'error')
        return redirect(request.referrer or url_for('tasks.index'))

    try:
        fname = secure_filename(file.filename)
        save_dir = Config.ACTIVITY_LOG_UPLOAD_DIR
        os.makedirs(save_dir, exist_ok=True)
        save_path = os.path.join(save_dir, fname)
        file.save(save_path)

        task_file = ActivityLogFile(
            log_id=task_id,
            file_name=fname,
            file_path=os.path.relpath(save_path, Config.UPLOAD_DIR),
            created_by=session.get('user_id', 0),
            created_at=datetime.utcnow()
        )
        db.session.add(task_file)
        db.session.commit()

        # If task has Jira ticket, upload attachment to Jira
        task = ActivityLog.query.get(task_id)
        if task and task.mt_jira_ticket:
            JiraService.upload_attachment(task.id, task.mt_jira_ticket, save_path, fname)

        flash('File successfully uploaded.', 'success')
    except Exception as e:
        flash(f'An error occurred: {str(e)}', 'error')

    return redirect(url_for('tasks.view', id=task_id))


@tasks_bp.route('/welcome-message')
@login_required
def welcome_message():
    user = session.get('user_name', 'there')
    return jsonify({'message': f"Hello {user}! Do you have any questions regarding the tasks? 😄"})


@tasks_bp.route('/comments/<int:log_id>')
@login_required
def get_comments(log_id):
    comments = ActivityLogComment.query.filter_by(log_id=log_id).order_by(ActivityLogComment.created_at.asc()).all()
    out = []
    for c in comments:
        out.append({
            'id': c.id,
            'comment': c.comment,
            'commented_by': c.commented_by,
            'user_name': Constants.USERS.get(c.commented_by, Constants.USER_SYSTEM_LABEL),
            'created_at': c.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            'is_system': c.commented_by == Constants.USER_SYSTEM
        })
    return jsonify(out)


@tasks_bp.route('/project-comments/<path:project_name>')
@login_required
def get_project_comments(project_name):
    comments = ProjectComment.query.filter_by(
        project_name=project_name,
        deleted_at=None
    ).order_by(ProjectComment.created_at.asc()).all()

    out = []
    current_uid = session.get('user_id')
    is_admin = (session.get('user_role') == Constants.USER_ROLE_ADMIN)

    for c in comments:
        u = User.query.get(c.commented_by)
        user_label = u.name if u else Constants.USERS.get(c.commented_by, 'Unknown')
        out.append({
            'id': c.id,
            'comment': c.comment,
            'commented_by': c.commented_by,
            'user_label': user_label,
            'formatted_time': c.created_at.strftime('%b %d, %Y, %I:%M %p') if c.created_at else '',
            'can_delete': (current_uid == c.commented_by) or is_admin
        })

    return jsonify({
        'success': True,
        'comments': out,
        'csrf_hash': ''
    })


@tasks_bp.route('/project-comments', methods=['POST'])
@tasks_bp.route('/post-project-comment', methods=['POST'])
@login_required
def post_project_comment():
    p_name = request.form.get('project_name', '').strip()
    c_text = request.form.get('comment', '').strip()
    if p_name and c_text:
        uid = session.get('user_id', 0)
        pc = ProjectComment(
            project_name=p_name,
            comment=c_text,
            commented_by=uid,
            created_at=datetime.utcnow()
        )
        db.session.add(pc)
        db.session.commit()

        u = User.query.get(uid)
        user_label = u.name if u else Constants.USERS.get(uid, session.get('user_name', 'You'))

        return jsonify({
            'success': True,
            'message': 'Comment added successfully.',
            'comment': {
                'id': pc.id,
                'comment': pc.comment,
                'user_label': user_label,
                'formatted_time': pc.created_at.strftime('%b %d, %Y, %I:%M %p') if pc.created_at else 'Just now',
                'can_delete': True
            },
            'csrf_hash': ''
        })
    return jsonify({'success': False, 'message': 'Missing data.'}), 400


@tasks_bp.route('/project-comments/delete', methods=['POST'])
@login_required
def delete_project_comment():
    comment_id = request.form.get('id', type=int)
    if not comment_id:
        return jsonify({'success': False, 'message': 'Missing comment ID.'}), 400
    pc = ProjectComment.query.get(comment_id)
    if not pc or pc.deleted_at is not None:
        return jsonify({'success': False, 'message': 'Comment not found.'}), 404

    uid = session.get('user_id', 0)
    is_admin = (session.get('user_role') == Constants.USER_ROLE_ADMIN)
    if pc.commented_by != uid and not is_admin:
        return jsonify({'success': False, 'message': 'Permission denied.'}), 403

    pc.deleted_at = datetime.utcnow()
    db.session.commit()
    return jsonify({'success': True, 'message': 'Comment deleted successfully.', 'csrf_hash': ''})


@tasks_bp.route('/project-comments/restore', methods=['POST'])
@login_required
def restore_project_comment():
    comment_id = request.form.get('id', type=int)
    if not comment_id:
        return jsonify({'success': False, 'message': 'Missing comment ID.'}), 400
    pc = ProjectComment.query.get(comment_id)
    if not pc:
        return jsonify({'success': False, 'message': 'Comment not found.'}), 404

    pc.deleted_at = None
    db.session.commit()
    return jsonify({'success': True, 'message': 'Comment restored.', 'csrf_hash': ''})


@tasks_bp.route('/export')
@login_required
def export():
    tasks = ActivityLog.query.filter(ActivityLog.deleted_at.is_(None)).order_by(ActivityLog.id.desc()).all()
    stream = ExcelService.export_tasks(tasks)
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"Tasks_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )


@tasks_bp.route('/project-export/excel/<path:project_name>')
@login_required
def export_project_excel(project_name):
    import io
    project_name = project_name.strip()
    show_completed = request.args.get('show_completed') == '1'
    me_mode = (session.get('al_me_mode') == 1)

    query = ActivityLog.query.filter(
        ActivityLog.project == project_name,
        ActivityLog.deleted_at.is_(None)
    )
    if me_mode:
        query = query.filter(ActivityLog.assignee == session.get('user_id'))
    if not show_completed:
        query = query.filter(ActivityLog.status != Constants.ACTIVITY_LOGS_STATUS_COMPLETED)

    tasks = query.order_by(ActivityLog.created_at.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tasks Export"

    headers = ['Subsidiary', 'Task Title', 'Ticket ID', 'Priority', 'Assignee', 'Created Date', 'Status', 'Latest Update']
    ws.append(headers)

    header_font = Font(name='Arial', size=10, bold=True, color='000000')
    header_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin', color='D1D1D6'),
        right=Side(style='thin', color='D1D1D6'),
        top=Side(style='thin', color='D1D1D6'),
        bottom=Side(style='thin', color='D1D1D6')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = thin_border

    # Get latest comments for tasks
    task_ids = [t.id for t in tasks]
    latest_comments = {}
    if task_ids:
        subq = db.session.query(
            ActivityLogComment.log_id,
            func.max(ActivityLogComment.id).label('max_id')
        ).filter(ActivityLogComment.log_id.in_(task_ids)).group_by(ActivityLogComment.log_id).subquery()

        comm_rows = db.session.query(ActivityLogComment).join(
            subq, (ActivityLogComment.id == subq.c.max_id)
        ).all()
        for cr in comm_rows:
            clean_comm = re.sub(r'<[^>]+>', '', cr.comment or '').strip()
            latest_comments[cr.log_id] = clean_comm

    row_font = Font(name='Arial', size=9)
    for task in tasks:
        subsidiary = Constants.ACTIVITY_LOG_SUBSIDIARIES.get(task.subsidiary_id, 'Unknown')
        priority = Constants.PRIORITIES.get(task.priority, 'Unknown')
        assignee = Constants.USERS.get(task.assignee, Constants.USER_SYSTEM_LABEL)
        created_date = task.created_at.strftime('%d-%m-%Y') if task.created_at else ''
        status = Constants.ACTIVITY_LOGS_STATUS.get(task.status, 'Unknown')
        last_comment = latest_comments.get(task.id, '')

        row_values = [
            subsidiary,
            task.title,
            task.ticket_id or '',
            priority,
            assignee,
            created_date,
            status,
            last_comment
        ]
        ws.append(row_values)
        curr_row = ws.max_row
        for col_idx in range(1, len(row_values) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.font = row_font
            c.border = thin_border
            c.alignment = Alignment(vertical='top')

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    safe_name = re.sub(r'[^a-zA-Z0-9_\-]', '_', project_name)
    download_name = f"{safe_name}_Tasks_Export_{datetime.now().strftime('%Y-%m-%d')}.xlsx"

    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=download_name
    )


@tasks_bp.route('/project-export/pdf/<path:project_name>')
@login_required
def export_project_pdf(project_name):
    project_name = project_name.strip()
    show_completed = request.args.get('show_completed') == '1'
    me_mode = (session.get('al_me_mode') == 1)

    query = ActivityLog.query.filter(
        ActivityLog.project == project_name,
        ActivityLog.deleted_at.is_(None)
    )
    if me_mode:
        query = query.filter(ActivityLog.assignee == session.get('user_id'))
    if not show_completed:
        query = query.filter(ActivityLog.status != Constants.ACTIVITY_LOGS_STATUS_COMPLETED)

    tasks = query.order_by(ActivityLog.created_at.desc()).all()

    # Preload latest comment
    task_ids = [t.id for t in tasks]
    latest_comments = {}
    if task_ids:
        subq = db.session.query(
            ActivityLogComment.log_id,
            func.max(ActivityLogComment.id).label('max_id')
        ).filter(ActivityLogComment.log_id.in_(task_ids)).group_by(ActivityLogComment.log_id).subquery()

        comm_rows = db.session.query(ActivityLogComment).join(
            subq, (ActivityLogComment.id == subq.c.max_id)
        ).all()
        for cr in comm_rows:
            clean_comm = re.sub(r'<[^>]+>', '', cr.comment or '').strip()
            latest_comments[cr.log_id] = clean_comm

    for t in tasks:
        t.last_comment = latest_comments.get(t.id, '')

    hide_assignee = False
    return render_template(
        'tasks/project_pdf.html',
        project_name=project_name,
        tasks=tasks,
        hide_assignee=hide_assignee,
        now=datetime.now(),
        constants=Constants
    )


@tasks_bp.route('/chat-with-ai', methods=['POST'])
@tasks_bp.route('/ai-chat', methods=['POST'])
@login_required
def chat_with_ai():
    api_url = Config.FABRIX_API_URL.rstrip('/') + '/openapi/llm/chat/completions'
    api_key = Config.FABRIX_API_TOKEN
    fabrix_client = Config.FABRIX_CLIENT

    data = request.get_json(silent=True) or request.form.to_dict()
    model_id = data.get('modelId')
    message = data.get('message', '').strip()
    raw_conv_id = data.get('conversationId')
    conversation_id = int(raw_conv_id) if raw_conv_id and str(raw_conv_id).isdigit() else None

    if not model_id or not message:
        return jsonify({'status': 'error', 'message': 'Missing modelId or message'}), 400

    db_messages = []
    conversation = None
    if conversation_id:
        conversation = AiConversation.query.get(conversation_id)
        if conversation and conversation.messages:
            try:
                db_messages = json.loads(conversation.messages)
            except Exception:
                db_messages = []

    db_messages.append({'role': 'user', 'content': message})

    current_uid = session.get('user_id', 0)
    current_uname = session.get('user_name') or session.get('username') or 'User'
    current_username = session.get('username') or ''
    current_date = datetime.now().strftime('%Y-%m-%d')
    current_datetime = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    current_day = datetime.now().strftime('%A')

    # Get Schema Context
    schema_context = "Database schema for Task tables:\n\n"
    tables = ['activity_logs', 'activity_log_comments', 'users', 'project_comments', 'properties', 'subsidiaries', 'tags', 'vendors', 'jira_history', 'tag_requests']
    for t in tables:
        try:
            res = db.session.execute(text(f"SHOW CREATE TABLE {t}")).fetchone()
            if res and len(res) > 1:
                schema_context += f"{res[1]};\n\n"
        except Exception:
            pass

    schema_context += "Note: team_id, subsidiary_id, property_id, status, type are integers mapping to constants. Here are the constant mappings you MUST use to convert natural language (like 'Media Tagging' or 'Pending') to integer IDs:\n"
    schema_context += f"TEAMS (team_id): {json.dumps(Constants.TEAMS)}\n"
    schema_context += f"SUBSIDIARIES (subsidiary_id): {json.dumps(Constants.ACTIVITY_LOG_SUBSIDIARIES)}\n"
    schema_context += f"STATUS (tags.status): {json.dumps(Constants.STATUS)}\n"
    schema_context += f"TICKET TYPES (jira_history.type, tag_requests.type): {json.dumps(Constants.JIRA_TICKET_TYPES)}\n"
    schema_context += f"PRIORITIES (activity_logs.priority): {json.dumps(Constants.PRIORITIES)}\n"
    schema_context += f"ACTIVITY_TYPES (activity_logs.type): {json.dumps(Constants.ACTIVITY_TYPES)}\n"
    schema_context += f"ACTIVITY_SIZES (activity_logs.size): {json.dumps(Constants.ACTIVITY_SIZES)}\n"
    schema_context += f"ACTIVITY_LOGS_STATUS (activity_logs.status): {json.dumps(Constants.ACTIVITY_LOGS_STATUS)}\n"
    schema_context += f"USERS (users.id / activity_logs.assignee / activity_logs.created_by): {json.dumps(Constants.USERS)}\n"
    schema_context += "For project, property, tag, or vendor names (and info like about/reason), join with the activity_logs.project, properties, tags, or vendors tables instead of guessing IDs.\n\n"

    user_context = "CURRENT AUTHENTICATED USER CONTEXT:\n"
    user_context += f"- Current User ID: {current_uid}\n"
    user_context += f"- Current User Name: {current_uname}\n"
    user_context += f"- Current User Username: {current_username}\n"
    user_context += f"- Today's Date (YYYY-MM-DD): {current_date} ({current_day})\n"
    user_context += f"- Current Timestamp: {current_datetime}\n\n"
    user_context += "CRITICAL TASK ASSIGNMENT & USER RULES:\n"
    user_context += "- In `activity_logs`, the `assignee` column is the integer ID of the assigned user, which maps to `users.id`.\n"
    user_context += "- In `activity_logs`, the `created_by` column is the integer ID of the creator user, which maps to `users.id`.\n"
    user_context += f"- When the user asks about their own tasks (using first-person words such as 'I', 'me', 'my tasks', 'assigned to me', 'tasks I completed', 'my pending tasks', 'what do I have to do'), you MUST filter by `activity_logs.assignee = {current_uid}`.\n"
    user_context += f"- For completed tasks, `activity_logs.status = {Constants.ACTIVITY_LOGS_STATUS_COMPLETED}`.\n"
    user_context += f"- When filtering tasks completed 'today', check `(DATE(activity_logs.finished_date) = '{current_date}' OR (activity_logs.status = {Constants.ACTIVITY_LOGS_STATUS_COMPLETED} AND DATE(activity_logs.updated_at) = '{current_date}'))`.\n"
    user_context += "- Exclude soft-deleted tasks by adding `activity_logs.deleted_at IS NULL` unless explicitly requested.\n\n"

    sql_prompt = [
        {'role': 'system', 'content': "You are an AI assistant that converts natural language requests into SQL queries. You will be provided with the database schema and current user context. Generate ONLY the raw SQL SELECT query required to fetch the data to answer the user's question. Do not add markdown formatting, do not add explanation, just the SQL query starting with SELECT. The query must be read-only (SELECT only). Limit results to a reasonable amount if necessary (e.g., LIMIT 50).\n\n" + user_context + schema_context},
        {'role': 'user', 'content': message}
    ]

    headers = {
        'Accept': 'application/json',
        'Content-Type': 'application/json',
        'x-fabrix-client': fabrix_client,
        'x-openapi-token': f"Bearer {api_key}",
        'x-llm-model-id': model_id
    }

    try:
        resp1 = requests.post(api_url, json={'messages': sql_prompt}, headers=headers, timeout=30)
        sql = ''
        if 200 <= resp1.status_code < 300:
            res_data1 = resp1.json()
            if 'choices' in res_data1 and res_data1['choices'] and 'message' in res_data1['choices'][0]:
                sql = res_data1['choices'][0]['message'].get('content', '').strip()
            elif 'reply' in res_data1:
                sql = res_data1.get('reply', '').strip()

        sql = re.sub(r'^```sql\s*', '', sql, flags=re.IGNORECASE)
        sql = re.sub(r'```$', '', sql).strip()

        data_result = []
        query_error = None
        if sql and sql.upper().startswith('SELECT'):
            try:
                raw_rows = db.session.execute(text(sql)).fetchall()
                data_result = [dict(row._mapping) for row in raw_rows]
            except Exception as e:
                query_error = str(e)
        else:
            query_error = f"Invalid or non-SELECT query generated by AI. Query was: {sql}"

        data_context = f"SQL Query Executed: {sql or 'None'}\n\n"
        if query_error:
            data_context += f"Error executing query: {query_error}"
        else:
            data_context += f"Query Result (JSON): \n{json.dumps(data_result, default=str)}"

        final_system_prompt = (
            "You are a helpful project management AI assistant. You answer questions about tasks.\n"
            f"You are currently assisting {current_uname} (User ID: {current_uid}, Username: {current_username}). Today is {current_day}, {current_date}.\n"
            "In `activity_logs`, the `assignee` column represents the assigned user's ID (mapping to `users.id`).\n"
            "You have executed a database query to find the answer. The query and the result data are provided below. Analyze the data and answer the user's question clearly and concisely based on their assigned tasks and the retrieved data.\n\n"
            + data_context
        )

        final_messages = [{'role': 'system', 'content': final_system_prompt}]
        final_messages.extend(db_messages)

        resp2 = requests.post(api_url, json={'messages': final_messages}, headers=headers, timeout=30)
        if 200 <= resp2.status_code < 300:
            res_data2 = resp2.json()
            reply = ''
            if 'choices' in res_data2 and res_data2['choices'] and 'message' in res_data2['choices'][0]:
                reply = res_data2['choices'][0]['message'].get('content', '')
            elif 'reply' in res_data2:
                reply = res_data2.get('reply', '')

            db_messages.append({'role': 'assistant', 'content': reply})

            if conversation:
                conversation.messages = json.dumps(db_messages)
                conversation.updated_at = datetime.utcnow()
                db.session.commit()
            else:
                title = message[:40] + ('...' if len(message) > 40 else '')
                conversation = AiConversation(
                    user_id=current_uid,
                    tag_id=0,
                    title=title,
                    messages=json.dumps(db_messages),
                    created_at=datetime.utcnow()
                )
                db.session.add(conversation)
                db.session.commit()
                conversation_id = conversation.id

            res_data2['conversationId'] = conversation_id
            return jsonify(res_data2)
        else:
            return jsonify({'status': 'error', 'message': f"API returned {resp2.status_code}"}), resp2.status_code

    except Exception as e:
        return jsonify({'status': 'error', 'message': str(e)}), 500


@tasks_bp.route('/ai-chat-history')
@tasks_bp.route('/ai-history')
@tasks_bp.route('/ai-history/<int:id>')
@login_required
def get_ai_chat_history(id=None):
    uid = session.get('user_id', 0)
    query = AiConversation.query.filter_by(user_id=uid)
    if id is not None:
        query = query.filter_by(tag_id=id)
    else:
        query = query.filter_by(tag_id=0)
    convs = query.order_by(AiConversation.updated_at.desc()).all()
    data = [{
        'id': c.id,
        'title': c.title,
        'tag_id': c.tag_id,
        'updated_at': c.updated_at.strftime('%Y-%m-%d %H:%M:%S') if c.updated_at else ''
    } for c in convs]
    return jsonify({'status': 'success', 'data': data})


@tasks_bp.route('/ai-conversation/<int:id>')
@login_required
def get_ai_conversation(id):
    c = AiConversation.query.get(id)
    if not c:
        return jsonify({'status': 'error', 'message': 'Conversation not found'}), 404
    return jsonify({
        'status': 'success',
        'data': {
            'id': c.id,
            'title': c.title,
            'messages': c.messages,
            'created_at': c.created_at.strftime('%Y-%m-%d %H:%M:%S') if c.created_at else ''
        }
    })

