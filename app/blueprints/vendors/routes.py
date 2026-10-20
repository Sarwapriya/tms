import os
import math
from io import BytesIO
from datetime import datetime, date, timedelta
from collections import defaultdict
from flask import Blueprint, render_template, request, redirect, url_for, session, flash, send_file
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models import db, Vendor, Tag, Property, Subsidiary
from app.constants import Constants
from app.utils.decorators import login_required, permission_required
from app.utils.helpers import format_date
from app.services.excel_service import ExcelService

vendors_bp = Blueprint('vendors', __name__)

@vendors_bp.route('/')
@login_required
@permission_required('vendors')
def index():
    page = request.args.get('page', 1, type=int)
    per_page = 25
    keyword = request.args.get('keyword', '').strip()
    vendor_id = request.args.get('id', type=int)
    sort_by = request.args.get('sort_by') or 'name'
    sort_dir = request.args.get('sort_dir', 'asc').lower()
    if sort_dir not in ('asc', 'desc'):
        sort_dir = 'asc'

    query = Vendor.query.filter_by(deleted_at=None)

    if keyword:
        if keyword.isdigit():
            query = query.filter((Vendor.id == int(keyword)) | (Vendor.name.ilike(f"%{keyword}%")))
        else:
            query = query.filter(Vendor.name.ilike(f"%{keyword}%"))

    if vendor_id:
        query = query.filter(Vendor.id == vendor_id)

    if sort_by == 'id':
        query = query.order_by(Vendor.id.desc() if sort_dir == 'desc' else Vendor.id.asc())
    else:
        query = query.order_by(Vendor.name.desc() if sort_dir == 'desc' else Vendor.name.asc())

    pagination = query.paginate(page=page, per_page=per_page, error_out=False)
    vendors = pagination.items

    today = date.today()
    expiring_date = today + timedelta(days=15)

    all_tags = db.session.query(
        Tag.id, Tag.vendor_id, Tag.status, Tag.al_end_date, Property.subsidiary_id
    ).join(Property, Tag.property_id == Property.id).filter(
        Tag.deleted_at.is_(None),
        Tag.status.notin_([
            Constants.STATUS_DISABLED,
            Constants.STATUS_ENDED,
            Constants.STATUS_EXPIRED
        ])
    ).all()

    grouped_tags = defaultdict(lambda: defaultdict(list))
    for t_id, v_id, status, end_date_val, sub_id in all_tags:
        if v_id and sub_id:
            grouped_tags[v_id][sub_id].append((status, end_date_val))

    vendor_matrix = defaultdict(dict)
    for v_id, subs in grouped_tags.items():
        for s_id, tag_list in subs.items():
            count = len(tag_list)
            has_expiring = False
            has_pending = False
            for st, ed in tag_list:
                if st in [Constants.STATUS_ENABLED, Constants.STATUS_PENDING_VERIFICATION, Constants.STATUS_UPDATED_IN_AL]:
                    if ed:
                        ed_date = ed.date() if isinstance(ed, datetime) else ed
                        if ed_date <= expiring_date:
                            has_expiring = True
                if st in [Constants.STATUS_PENDING_VERIFICATION, Constants.STATUS_NEW_TAG_SUBMITTED, Constants.STATUS_NEW_RULE_SUBMITTED]:
                    has_pending = True
            if has_expiring:
                badge_class = 'bubble-red'
            elif has_pending:
                badge_class = 'bubble-amber'
            else:
                badge_class = 'bubble-green'
            vendor_matrix[v_id][s_id] = {'count': count, 'badge_class': badge_class}

    advanced_filter_keys = ['id']
    active_filters_count = 1 if (vendor_id is not None and vendor_id != '') else 0
    if keyword:
        active_filters_count += 1
    has_active_filters = (active_filters_count > 0)

    filters = {
        'keyword': keyword,
        'id': vendor_id,
        'sort_by': sort_by,
        'sort_dir': sort_dir
    }

    return render_template(
        'vendors/list.html',
        vendors=vendors,
        pagination=pagination,
        pager=pagination,
        totalVendors=pagination.total,
        vendorOptions=Vendor.options(),
        vendor_matrix=vendor_matrix,
        activeFiltersCount=active_filters_count,
        hasActiveFilters=has_active_filters,
        advancedFilterKeys=advanced_filter_keys,
        filters=filters,
        Constants=Constants
    )

@vendors_bp.route('/<int:id>', endpoint='view_short')
@vendors_bp.route('/view/<int:id>')
@login_required
@permission_required('vendors')
def view(id):
    vendor = Vendor.query.get_or_404(id)
    subsidiary_id = request.args.get('subsidiary_id', type=int)
    subsidiary = Subsidiary.query.get(subsidiary_id) if subsidiary_id else Subsidiary.query.first()
    if not subsidiary:
        subsidiary = Subsidiary(id=0, name='Global')
    query = Tag.query.filter_by(vendor_id=vendor.id, deleted_at=None)
    if subsidiary_id:
        query = query.join(Property).filter(Property.subsidiary_id == subsidiary_id)
    tags = query.order_by(Tag.id.desc()).all()

    properties_map = {p.id: p.name for p in Property.query.all()}

    return render_template(
        'vendors/view.html',
        vendor=vendor,
        tags=tags,
        subsidiary=subsidiary,
        subsidiary_id=subsidiary.id,
        properties=properties_map
    )

@vendors_bp.route('/create', methods=['GET', 'POST'])
@login_required
@permission_required('vendors')
def create():
    if request.method == 'POST':
        name = request.form.get('name')
        about = request.form.get('about')
        reason = request.form.get('reason')
        tag_purpose = request.form.get('tag_purpose')
        collected_data = request.form.get('collected_data')

        vendor = Vendor(
            name=name,
            about=about,
            reason=reason,
            tag_purpose=tag_purpose,
            collected_data=collected_data,
            created_by=session.get('user_id', 0)
        )
        db.session.add(vendor)
        db.session.commit()
        flash(f"Vendor '{vendor.name}' added successfully.", 'success')
        return redirect(url_for('vendors.index'))

    return render_template('vendors/create.html')

@vendors_bp.route('/update/<int:id>', methods=['GET', 'POST'])
@login_required
@permission_required('vendors')
def update(id):
    vendor = Vendor.query.get_or_404(id)
    subsidiary_id = request.args.get('subsidiary_id', type=int) or request.form.get('subsidiary_id', type=int)

    if request.method == 'POST':
        name = (request.form.get('name') or '').strip()
        vendor.name = name
        vendor.about = request.form.get('about')
        vendor.reason = request.form.get('reason')
        vendor.tag_purpose = request.form.get('tag_purpose')
        vendor.collected_data = request.form.get('collected_data')
        vendor.updated_by = session.get('user_id', 0)
        db.session.commit()
        flash(f"Vendor '{vendor.name}' updated successfully.", 'success')
        if subsidiary_id:
            return redirect(url_for('vendors.view', id=vendor.id, subsidiary_id=subsidiary_id))
        return redirect(url_for('vendors.view', id=vendor.id))

    return render_template('vendors/update.html', vendor=vendor, subsidiary_id=subsidiary_id)

@vendors_bp.route('/export')
@vendors_bp.route('/export/<int:vendor_id>')
@vendors_bp.route('/export-excel')
@vendors_bp.route('/export-excel/<int:vendor_id>')
@login_required
@permission_required('vendors')
def export(vendor_id=None):
    wb = openpyxl.Workbook()
    ws = wb.active

    thin_side = Side(border_style="thin", color="CCCCCC")
    table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True)
    data_font = Font(name="Calibri", size=10)

    # 1. Single Vendor Export with subsidiary filter
    if vendor_id is not None:
        vendor = Vendor.query.get_or_404(vendor_id)
        subsidiary_id = request.args.get('subsidiary_id', type=int)
        subsidiary = Subsidiary.query.get(subsidiary_id) if subsidiary_id else None

        q = Tag.query.filter_by(vendor_id=vendor.id, deleted_at=None)
        if subsidiary_id:
            q = q.join(Property).filter(Property.subsidiary_id == subsidiary_id)
        tags = q.order_by(Tag.id.desc()).all()

        ws.title = "Vendor Tags"
        headers = ['Subsidiary', 'Property Name', 'Tag Name', 'Days', 'AL End Date', 'Requested Date', 'Updated', 'Status']
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.font = header_font
            c.fill = header_fill
            c.border = table_border
            c.alignment = Alignment(horizontal="center" if col_idx in [1, 4, 5, 6, 7, 8] else "left")

        today_dt = datetime.now()
        for tag in tags:
            prop = tag.property
            sub_name = subsidiary.name if subsidiary else (prop.subsidiary.name if prop and prop.subsidiary else 'N/A')
            prop_name = prop.name if prop else 'N/A'

            # Days calculation matching PHP ceil((strtotime(al_end_date) - time()) / 86400)
            days_str = 'N/A'
            if tag.al_end_date:
                try:
                    ed_datetime = datetime.combine(tag.al_end_date, datetime.min.time()) if isinstance(tag.al_end_date, date) and not isinstance(tag.al_end_date, datetime) else tag.al_end_date
                    diff_seconds = (ed_datetime - today_dt).total_seconds()
                    days_str = str(math.ceil(diff_seconds / 86400.0))
                except Exception:
                    pass

            row = [
                sub_name,
                prop_name,
                tag.name,
                days_str,
                format_date(tag.al_end_date),
                format_date(tag.requested_date or tag.al_end_date) or '-',
                format_date(tag.updated_at) or '-',
                Constants.STATUS.get(tag.status, 'Unknown')
            ]
            ws.append(row)
            curr_r = ws.max_row
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_r, column=col_idx)
                cell.font = data_font
                cell.border = table_border

        ExcelService.auto_fit_columns(ws)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        sub_label = subsidiary.name if subsidiary else 'All'
        filename = f"Tags_[{sub_label}][{vendor.name}]_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return send_file(
            stream,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename
        )

    # 2. Export All Vendors List (Subsidiary matrix format matching PHP)
    keyword = request.args.get('keyword', '').strip()
    vendor_filter_id = request.args.get('id', type=int)
    sort_by = request.args.get('sort_by') or 'name'
    sort_dir = request.args.get('sort_dir', 'asc').lower()

    query = Vendor.query.filter_by(deleted_at=None)
    if keyword:
        if keyword.isdigit():
            query = query.filter((Vendor.id == int(keyword)) | (Vendor.name.ilike(f"%{keyword}%")))
        else:
            query = query.filter(Vendor.name.ilike(f"%{keyword}%"))
    if vendor_filter_id:
        query = query.filter(Vendor.id == vendor_filter_id)

    if sort_by == 'id':
        query = query.order_by(Vendor.id.desc() if sort_dir == 'desc' else Vendor.id.asc())
    else:
        query = query.order_by(Vendor.name.desc() if sort_dir == 'desc' else Vendor.name.asc())

    vendors = query.all()

    all_tags = db.session.query(
        Tag.id, Tag.vendor_id, Tag.status, Property.subsidiary_id
    ).join(Property, Tag.property_id == Property.id).filter(
        Tag.deleted_at.is_(None),
        Tag.status.notin_([
            Constants.STATUS_DISABLED,
            Constants.STATUS_ENDED,
            Constants.STATUS_EXPIRED
        ])
    ).all()

    tags_by_vendor_and_sub = defaultdict(lambda: defaultdict(int))
    for t_id, v_id, status, sub_id in all_tags:
        if v_id and sub_id:
            tags_by_vendor_and_sub[v_id][sub_id] += 1

    ws.title = "Vendors"
    headers = ['#', 'Vendor']
    subs_list = list(Constants.SUBSIDIARIES.items())
    for sub_id, sub_name in subs_list:
        headers.append(sub_name)
    headers.append('Total')

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.border = table_border
        c.alignment = Alignment(horizontal="center" if col_idx != 2 else "left")

    for v in vendors:
        row = [f"#{v.id}", v.name]
        vendor_total = 0
        for sub_id, sub_name in subs_list:
            count = tags_by_vendor_and_sub[v.id][sub_id]
            vendor_total += count
            row.append(count if count > 0 else '-')
        row.append(vendor_total if vendor_total > 0 else '-')
        ws.append(row)

        curr_r = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_r, column=col_idx)
            cell.font = data_font
            cell.border = table_border
            cell.alignment = Alignment(horizontal="center" if col_idx != 2 else "left")
            if col_idx == len(headers):
                cell.font = Font(name="Calibri", size=10, bold=True)

    ExcelService.auto_fit_columns(ws)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    filename = f"Vendors_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
