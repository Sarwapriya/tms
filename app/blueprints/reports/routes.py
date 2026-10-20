import os
from io import BytesIO
from datetime import datetime, date, timedelta
from collections import defaultdict
from flask import Blueprint, render_template, request, session, send_file, redirect, url_for
from sqlalchemy import or_, and_, func
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.models import db, ActivityLog, User, Subsidiary, Tag, TagRequest, Property, Vendor
from app.constants import Constants
from app.utils.decorators import login_required, permission_required
from app.utils.helpers import get_system_settings
from app.services.excel_service import ExcelService

reports_bp = Blueprint('reports', __name__)

def get_report_data():
    today = date.today()
    week = request.args.get('week')
    start_param = request.args.get('start_date')
    end_param = request.args.get('end_date')

    if not week and not start_param:
        week = today.strftime('%Y-W%W')

    start_date = None
    end_date = None

    if week:
        try:
            parts = week.split('-W')
            if len(parts) == 2:
                year = int(parts[0])
                week_num = int(parts[1])
                start_date = date.fromisocalendar(year, week_num, 1)
                end_date = start_date + timedelta(days=6)
        except Exception:
            pass

    if not start_date:
        try:
            start_date = datetime.strptime(start_param, '%Y-%m-%d').date() if start_param else (today - timedelta(days=today.weekday()))
            end_date = datetime.strptime(end_param, '%Y-%m-%d').date() if end_param else (start_date + timedelta(days=6))
        except Exception:
            start_date = today - timedelta(days=6)
            end_date = today

    if start_date and end_date and start_date > end_date:
        start_date = end_date

    start_dt = datetime.combine(start_date, datetime.min.time())
    end_dt = datetime.combine(end_date, datetime.max.time())

    # User Role and Team restrictions
    user_id = session.get('user_id')
    user_role = session.get('user_role')
    team_param = request.args.get('team_id')
    selected_team = None
    if team_param is not None and str(team_param).strip() != '':
        try:
            selected_team = int(team_param)
        except Exception:
            pass
    elif user_role != Constants.USER_ROLE_ADMIN:
        selected_team = Constants.USER_TEAMS.get(user_id)

    # 1. Completed tasks (Finished within range)
    completed_q = ActivityLog.query.filter(
        ActivityLog.status == Constants.ACTIVITY_LOGS_STATUS_COMPLETED,
        ActivityLog.finished_date >= start_date,
        ActivityLog.finished_date <= end_date,
        ActivityLog.deleted_at.is_(None)
    )
    if selected_team is not None:
        completed_q = completed_q.filter(ActivityLog.team_id == selected_team)
    completed_tasks = completed_q.order_by(ActivityLog.finished_date.desc()).all()

    # 2. New tasks created (Created within range)
    new_q = ActivityLog.query.filter(
        ActivityLog.created_at >= start_dt,
        ActivityLog.created_at <= end_dt,
        ActivityLog.deleted_at.is_(None)
    )
    if selected_team is not None:
        new_q = new_q.filter(ActivityLog.team_id == selected_team)
    new_tasks = new_q.order_by(ActivityLog.created_at.desc()).all()

    # 3. Ongoing tasks (Created on/before end date and either not completed or completed after end date)
    ongoing_q = ActivityLog.query.filter(
        ActivityLog.deleted_at.is_(None),
        func.date(ActivityLog.created_at) <= end_date,
        or_(
            ActivityLog.status != Constants.ACTIVITY_LOGS_STATUS_COMPLETED,
            and_(
                ActivityLog.status == Constants.ACTIVITY_LOGS_STATUS_COMPLETED,
                ActivityLog.finished_date > end_date
            )
        )
    )
    if selected_team is not None:
        ongoing_q = ongoing_q.filter(ActivityLog.team_id == selected_team)
    ongoing_tasks = ongoing_q.order_by(ActivityLog.priority.desc(), ActivityLog.created_at.desc()).all()

    # 4. Tag requests & tags (Media Tagging team only)
    if selected_team is not None and selected_team != Constants.TEAM_MEDIA_TAGGING:
        new_requests = []
        published_requests = []
        new_tags = []
        updated_tags = []
    else:
        new_requests = TagRequest.query.filter(
            TagRequest.created_at >= start_dt,
            TagRequest.created_at <= end_dt,
            TagRequest.deleted_at.is_(None)
        ).order_by(TagRequest.created_at.desc()).all()

        published_requests = TagRequest.query.filter(
            TagRequest.status == Constants.STATUS_REQUEST_PUBLISHED,
            TagRequest.updated_at >= start_dt,
            TagRequest.updated_at <= end_dt,
            TagRequest.deleted_at.is_(None)
        ).order_by(TagRequest.updated_at.desc()).all()

        new_tags = Tag.query.filter(
            Tag.created_at >= start_dt,
            Tag.created_at <= end_dt,
            Tag.deleted_at.is_(None)
        ).order_by(Tag.created_at.desc()).all()

        updated_tags = Tag.query.filter(
            Tag.updated_at >= start_dt,
            Tag.updated_at <= end_dt,
            Tag.created_at < start_dt,
            Tag.deleted_at.is_(None)
        ).order_by(Tag.updated_at.desc()).all()

    # Helper mappings
    props_cache = {p.id: p for p in Property.query.all()}
    subs_cache = {s.id: s.name for s in Subsidiary.query.all()}
    vendors_cache = {v.id: v.name for v in Vendor.query.all()}
    users_cache = {u.id: u.name for u in User.query.all()}

    for item in completed_tasks + new_tasks + ongoing_tasks:
        sub_id = item.subsidiary_id or 0
        item.subsidiary_name = subs_cache.get(sub_id, Constants.ACTIVITY_LOG_SUBSIDIARIES.get(sub_id, 'Global'))
        item.property_name = getattr(item, 'property_name', 'N/A')
        item.assignee_name = users_cache.get(item.assignee, Constants.USERS.get(item.assignee, Constants.USER_SYSTEM_LABEL))

    for r in new_requests + published_requests:
        prop = props_cache.get(r.property_id)
        r.property_name = prop.name if prop else 'N/A'
        r.vendor_name = vendors_cache.get(r.vendor_id, 'Unknown')
        sub_id = prop.subsidiary_id if prop else 0
        r.subsidiary_name = subs_cache.get(sub_id, Constants.ACTIVITY_LOG_SUBSIDIARIES.get(sub_id, 'Global'))

    for t in new_tags + updated_tags:
        prop = props_cache.get(t.property_id)
        t.property_name = prop.name if prop else 'N/A'
        t.vendor_name = vendors_cache.get(t.vendor_id, 'Unknown')
        sub_id = prop.subsidiary_id if prop else 0
        t.subsidiary_name = subs_cache.get(sub_id, Constants.ACTIVITY_LOG_SUBSIDIARIES.get(sub_id, 'Global'))

    # Status counts for chart
    status_counts = {
        'To Do': sum(1 for t in ongoing_tasks if t.status == Constants.ACTIVITY_LOGS_STATUS_TO_DO),
        'In Progress': sum(1 for t in ongoing_tasks if t.status in [Constants.ACTIVITY_LOGS_STATUS_INVESTIGATION, Constants.ACTIVITY_LOGS_STATUS_IMPLEMENTATION]),
        'QA/Review': sum(1 for t in ongoing_tasks if t.status in [Constants.ACTIVITY_LOGS_STATUS_QA, Constants.ACTIVITY_LOGS_STATUS_REVIEW]),
        'Completed': len(completed_tasks),
        'Blocked': sum(1 for t in ongoing_tasks if t.status == Constants.ACTIVITY_LOGS_STATUS_BLOCKED)
    }

    subsidiary_activity = defaultdict(int)
    for s_name in Constants.ACTIVITY_LOG_SUBSIDIARIES.values():
        subsidiary_activity[s_name] = 0
    for t in new_tasks + completed_tasks:
        sub_name = subs_cache.get(t.subsidiary_id, Constants.ACTIVITY_LOG_SUBSIDIARIES.get(t.subsidiary_id))
        if sub_name:
            subsidiary_activity[sub_name] += 1
    for t in new_tags:
        prop = props_cache.get(t.property_id)
        sub_id = prop.subsidiary_id if prop else 0
        sub_name = subs_cache.get(sub_id, Constants.ACTIVITY_LOG_SUBSIDIARIES.get(sub_id))
        if sub_name:
            subsidiary_activity[sub_name] += 1
    for r in new_requests:
        prop = props_cache.get(r.property_id)
        sub_id = prop.subsidiary_id if prop else 0
        sub_name = subs_cache.get(sub_id, Constants.ACTIVITY_LOG_SUBSIDIARIES.get(sub_id))
        if sub_name:
            subsidiary_activity[sub_name] += 1

    project_activity = defaultdict(int)
    for t in new_tasks + completed_tasks:
        p_name = t.project or 'Generic'
        project_activity[p_name] += 1

    assignee_workload = defaultdict(int)
    for t in completed_tasks:
        u_name = users_cache.get(t.assignee, Constants.USERS.get(t.assignee, Constants.USER_SYSTEM_LABEL))
        assignee_workload[u_name] += 1

    # Check system setting for hideAssignee
    sys_settings = get_system_settings()
    hide_assignee = bool(sys_settings.get('hide_assignee_system_wide') or sys_settings.get('hide_assignee_weekly_reports'))

    return {
        'week': week or '',
        'start_date': start_date.strftime('%Y-%m-%d'),
        'end_date': end_date.strftime('%Y-%m-%d'),
        'team_id': selected_team,
        'completed_tasks': completed_tasks,
        'new_tasks': new_tasks,
        'ongoing_tasks': ongoing_tasks,
        'new_requests': new_requests,
        'published_requests': published_requests,
        'new_tags': new_tags,
        'updated_tags': updated_tags,
        'status_counts': status_counts,
        'subsidiary_activity': dict(subsidiary_activity),
        'project_activity': dict(project_activity),
        'assignee_workload': dict(assignee_workload),
        'hideAssignee': hide_assignee,
        'summaryParams': {'summary': '1', 'week': week, 'start_date': start_date.strftime('%Y-%m-%d'), 'end_date': end_date.strftime('%Y-%m-%d')},
        'detailedParams': {'detailed': '1', 'week': week, 'start_date': start_date.strftime('%Y-%m-%d'), 'end_date': end_date.strftime('%Y-%m-%d')},
        'inventoryParams': {'inventory': '1', 'week': week, 'start_date': start_date.strftime('%Y-%m-%d'), 'end_date': end_date.strftime('%Y-%m-%d')}
    }

@reports_bp.route('/')
@login_required
@permission_required('reports')
def index():
    data = get_report_data()
    return render_template('reports/index.html', **data)

@reports_bp.route('/pdf')
@reports_bp.route('/export/pdf')
@login_required
@permission_required('reports')
def export_pdf():
    data = get_report_data()
    report_type = 'detailed'
    if request.args.get('summary') == '1':
        report_type = 'summary'
    elif request.args.get('inventory') == '1':
        report_type = 'inventory'
    data['reportType'] = report_type
    data['report_type'] = report_type
    data['prioClass'] = lambda p: Constants.PRIORITIES_COLOR.get(p, 'secondary')
    data['statusClass'] = lambda s: Constants.ACTIVITY_LOGS_STATUS_COLOR.get(s, 'secondary')
    return render_template('reports/pdf.html', **data)

@reports_bp.route('/export/excel')
@reports_bp.route('/export-excel')
@login_required
@permission_required('reports')
def export_excel():
    data = get_report_data()
    hide_assignee = data.get('hideAssignee', False)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Report"

    # Styling definitions
    title_font = Font(name="Calibri", size=14, bold=True, color="1F2937")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="4B5563")
    section_font = Font(name="Calibri", size=12, bold=True, color="1F2937")
    header_font = Font(name="Calibri", size=10, bold=True, color="000000")
    data_font = Font(name="Calibri", size=10, color="000000")
    header_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    kpi_metric_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CCCCCC")
    table_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # 1. Main Title
    ws.append(["Weekly Performance & Activity Report"])
    ws.cell(row=1, column=1).font = title_font

    team_str = f" | Team: {Constants.TEAMS.get(data['team_id'], 'Unknown')}" if data.get('team_id') else ""
    ws.append([f"Period: {data['start_date']} to {data['end_date']}{team_str} | Generated At: {datetime.now().strftime('%d-%m-%Y %H:%M:%S')}"])
    ws.cell(row=2, column=1).font = subtitle_font
    ws.append([])  # Blank row

    # Helper function to append styled table
    def add_table(title, headers, rows, empty_msg):
        # Section header
        r = ws.max_row + 1
        ws.cell(row=r, column=1, value=title).font = section_font
        
        # Table Header
        ws.append(headers)
        h_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=h_row, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = table_border
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3, 4, len(headers)] else "left")

        # Table Rows
        if not rows:
            ws.append([empty_msg] + [''] * (len(headers) - 1))
            row_idx = ws.max_row
            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(headers))
            c = ws.cell(row=row_idx, column=1)
            c.font = Font(name="Calibri", size=10, italic=True, color="6B7280")
            c.border = table_border
            c.alignment = Alignment(horizontal="center")
        else:
            for row_data in rows:
                ws.append(row_data)
                curr_r = ws.max_row
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=curr_r, column=col_idx)
                    cell.font = data_font
                    cell.border = table_border
        ws.append([])  # Spacer row

    # Section 1: KPIs
    kpi_rows = [
        ["Completed Tasks", len(data['completed_tasks'])],
        ["New Tasks Created", len(data['new_tasks'])],
        ["New Tag Requests Submitted", len(data['new_requests'])],
        ["Tags Published", len(data['published_requests'])],
        ["Total Tags Created/Updated", len(data['new_tags']) + len(data['updated_tags'])],
    ]
    add_table("Key Metrics Summary", ["Metric", "Count"], kpi_rows, "No metrics available.")

    # Section 2: Completed Tasks
    comp_headers = ["Subsidiary", "Task Title", "Ticket ID", "Priority"]
    if not hide_assignee:
        comp_headers.append("Assignee")
    comp_headers.append("Finished Date")

    comp_rows = []
    for t in data['completed_tasks']:
        row = [
            t.subsidiary_name,
            t.title or '',
            t.ticket_id or '—',
            Constants.PRIORITIES.get(t.priority, 'N/A')
        ]
        if not hide_assignee:
            row.append(t.assignee_name)
        row.append(t.finished_date.strftime('%Y-%m-%d') if t.finished_date else '')
        comp_rows.append(row)
    add_table("Completed Tasks", comp_headers, comp_rows, "No completed tasks recorded in this period.")

    # Section 3: New Tag Requests Submitted
    req_headers = ["Ticket ID / Request", "Type", "Property", "Vendor", "Status", "Requested By", "Created Date"]
    req_rows = []
    for r in data['new_requests']:
        ticket = r.new_rule_ticket if r.type == Constants.JIRA_TYPE_NEW_RULE else r.new_tag_ticket
        req_rows.append([
            ticket or '—',
            "New Rule" if r.type == Constants.JIRA_TYPE_NEW_RULE else "New Tag",
            r.property_name,
            r.vendor_name,
            Constants.STATUS_REQUESTS.get(r.status, 'Unknown'),
            r.requester_name or '—',
            r.created_at.strftime('%Y-%m-%d %H:%M:%S') if r.created_at else ''
        ])
    add_table("New Tag Requests Submitted", req_headers, req_rows, "No new tag requests submitted in this period.")

    # Section 4: Tags Published
    pub_headers = ["Ticket ID", "Type", "Property", "Vendor", "Requester", "Published Date"]
    pub_rows = []
    for r in data['published_requests']:
        ticket = r.new_rule_ticket if r.type == Constants.JIRA_TYPE_NEW_RULE else r.new_tag_ticket
        pub_rows.append([
            ticket or '—',
            "New Rule" if r.type == Constants.JIRA_TYPE_NEW_RULE else "New Tag",
            r.property_name,
            r.vendor_name,
            r.requester_name or '—',
            r.updated_at.strftime('%Y-%m-%d %H:%M:%S') if r.updated_at else ''
        ])
    add_table("Tags Published", pub_headers, pub_rows, "No tags published in this period.")

    # Section 5: Ongoing Tasks
    ong_headers = ["Subsidiary", "Task Title", "Ticket ID", "Priority"]
    if not hide_assignee:
        ong_headers.append("Assignee")
    ong_headers.extend(["Status", "Created Date"])

    ong_rows = []
    for t in data['ongoing_tasks']:
        row = [
            t.subsidiary_name,
            t.title or '',
            t.ticket_id or '—',
            Constants.PRIORITIES.get(t.priority, 'N/A')
        ]
        if not hide_assignee:
            row.append(t.assignee_name)
        row.extend([
            Constants.ACTIVITY_LOGS_STATUS.get(t.status, 'N/A'),
            t.created_at.strftime('%Y-%m-%d %H:%M:%S') if t.created_at else ''
        ])
        ong_rows.append(row)
    add_table("Ongoing Tasks (Todo, In Progress, Blocked)", ong_headers, ong_rows, "No active ongoing tasks recorded.")

    ExcelService.auto_fit_columns(ws)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"Weekly_Report_{data['start_date']}_to_{data['end_date']}.xlsx"
    )
