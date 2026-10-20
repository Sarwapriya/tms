import os
from io import BytesIO
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

class ExcelService:
    @staticmethod
    def apply_header_style(ws, row_num=1, bg_color="0071E3", font_color="FFFFFF"):
        fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
        font = Font(name="Segoe UI", size=11, bold=True, color=font_color)
        alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='CCCCCC'),
            right=Side(style='thin', color='CCCCCC'),
            top=Side(style='thin', color='CCCCCC'),
            bottom=Side(style='thin', color='CCCCCC')
        )
        
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.fill = fill
            cell.font = font
            cell.alignment = alignment
            cell.border = thin_border
        ws.row_dimensions[row_num].height = 28

    @staticmethod
    def auto_fit_columns(ws, max_len_cap=60):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), max_len_cap)

    @staticmethod
    def export_tags(tags, selected_fields=None, user_map=None, org_map=None):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tags Inventory"

        from app.constants import Constants
        from app.models.user import User
        from app.models.metadata import Organization

        if user_map is None:
            user_map = {u.id: u.full_name for u in User.query.all()}
        if org_map is None:
            org_map = {o.id: o.name for o in Organization.query.all()}

        field_definitions = {
            'subsidiary': {
                'label': 'Subsidiary',
                'getter': lambda tag: (tag.property.subsidiary.name if tag.property and tag.property.subsidiary 
                                       else (Constants.SUBSIDIARIES.get(tag.property.subsidiary_id, '') if tag.property and getattr(tag.property, 'subsidiary_id', None) else ''))
            },
            'property_name': {
                'label': 'Property Name',
                'getter': lambda tag: tag.property.name if tag.property else ''
            },
            'tag_name': {
                'label': 'Tag Name',
                'getter': lambda tag: tag.name or ''
            },
            'days': {
                'label': 'Days',
                'getter': lambda tag: tag.days_remaining if tag.days_remaining is not None else ''
            },
            'al_end_date': {
                'label': 'AL End Date',
                'getter': lambda tag: tag.al_end_date.strftime('%Y-%m-%d') if tag.al_end_date else ''
            },
            'requested_date': {
                'label': 'Requested Date',
                'getter': lambda tag: tag.requested_date.strftime('%Y-%m-%d') if tag.requested_date else ''
            },
            'updated_at': {
                'label': 'Updated',
                'getter': lambda tag: tag.updated_at.strftime('%Y-%m-%d %H:%M:%S') if tag.updated_at else ''
            },
            'status': {
                'label': 'Status',
                'getter': lambda tag: Constants.TAG_STATUS.get(tag.status, str(tag.status))
            },
            'tag_id': {
                'label': 'Tag ID',
                'getter': lambda tag: tag.id
            },
            'requested_name': {
                'label': 'Requested Name',
                'getter': lambda tag: tag.requested_name or ''
            },
            'vendor': {
                'label': 'Vendor',
                'getter': lambda tag: tag.vendor.name if tag.vendor else ''
            },
            'conditions': {
                'label': 'Firing Conditions',
                'getter': lambda tag: tag.conditions or ''
            },
            'ref_link': {
                'label': 'AL Ref Link',
                'getter': lambda tag: tag.ref_link or ''
            },
            'requester_name': {
                'label': 'Requester Name',
                'getter': lambda tag: tag.requester_name or ''
            },
            'requester_email': {
                'label': 'Requester Email',
                'getter': lambda tag: tag.requester_email or ''
            },
            'requester_organization': {
                'label': 'Requester Organization',
                'getter': lambda tag: org_map.get(tag.requester_organization, str(tag.requester_organization or ''))
            },
            'js_code': {
                'label': 'JavaScript Code',
                'getter': lambda tag: tag.js_code or ''
            },
            'property_website_url': {
                'label': 'Property Website URL',
                'getter': lambda tag: (tag.property.website or '') if tag.property else ''
            },
            'created_at': {
                'label': 'Created Date',
                'getter': lambda tag: tag.created_at.strftime('%Y-%m-%d %H:%M:%S') if tag.created_at else ''
            },
            'created_by': {
                'label': 'Created By',
                'getter': lambda tag: user_map.get(tag.created_by, '') if tag.created_by else ''
            },
            'updated_by': {
                'label': 'Updated By',
                'getter': lambda tag: user_map.get(tag.updated_by, '') if tag.updated_by else ''
            },
        }

        default_field_keys = [
            'subsidiary', 'property_name', 'tag_name', 'days',
            'al_end_date', 'requested_date', 'updated_at', 'status'
        ]

        if not selected_fields:
            active_fields = default_field_keys
        else:
            active_fields = [f for f in selected_fields if f in field_definitions]
            if not active_fields:
                active_fields = default_field_keys

        headers = [field_definitions[f]['label'] for f in active_fields]
        ws.append(headers)

        for tag in tags:
            row = [field_definitions[f]['getter'](tag) for f in active_fields]
            ws.append(row)

        ExcelService.apply_header_style(ws, 1)
        ExcelService.auto_fit_columns(ws)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def export_tag_requests(requests_list):
        from app.constants import Constants
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tag Requests"

        headers = ['Subsidiary', 'Property Name', 'Ticket ID', 'Ticket Type', 'Status', 'Created At']
        ws.append(headers)

        for req in requests_list:
            prop = req.property if hasattr(req, 'property') and req.property else None
            sub_name = 'N/A'
            if prop and hasattr(prop, 'subsidiary') and prop.subsidiary:
                sub_name = prop.subsidiary.name
            elif prop and getattr(prop, 'subsidiary_id', None):
                sub_name = Constants.SUBSIDIARIES.get(prop.subsidiary_id, 'N/A')

            prop_name = prop.name if prop else 'N/A'
            ticket_id = (req.new_rule_ticket if req.type == Constants.JIRA_TYPE_NEW_RULE else req.new_tag_ticket) or 'N/A'
            type_label = 'New Rule' if req.type == Constants.JIRA_TYPE_NEW_RULE else 'New Tag'
            status_label = Constants.STATUS_REQUESTS.get(req.status, 'Unknown')
            created_at_str = req.created_at.strftime('%Y-%m-%d') if req.created_at else ''

            ws.append([
                sub_name,
                prop_name,
                ticket_id,
                type_label,
                status_label,
                created_at_str
            ])

        ExcelService.apply_header_style(ws, 1)
        ExcelService.auto_fit_columns(ws)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def export_tasks(tasks, title="Activity Logs"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tasks"

        headers = [
            "ID", "Title", "Project", "Ticket ID", "MT Jira Ticket",
            "Assignee", "Status", "Priority", "Size", "Requested Date", "Finished Date"
        ]
        ws.append(headers)

        for t in tasks:
            ws.append([
                t.id,
                t.title,
                t.project or '',
                t.ticket_id or '',
                t.mt_jira_ticket or '',
                t.assignee or '',
                t.status,
                t.priority,
                t.size,
                t.requested_date.strftime('%Y-%m-%d') if t.requested_date else '',
                t.finished_date.strftime('%Y-%m-%d') if t.finished_date else ''
            ])

        ExcelService.apply_header_style(ws, 1)
        ExcelService.auto_fit_columns(ws)

        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    @staticmethod
    def parse_excel_file(file_path):
        """Parses an uploaded Excel workbook and returns list of rows as dicts."""
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        
        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        data = []
        for r in rows[1:]:
            if any(cell is not None for cell in r):
                row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
                data.append(row_dict)
        return data
